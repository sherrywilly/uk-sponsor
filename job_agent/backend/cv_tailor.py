from __future__ import annotations

import json
import shutil
import subprocess
from datetime import datetime
from pathlib import Path
from typing import Any

try:
    from docx import Document
except ImportError:  # pragma: no cover
    Document = None

from .claude_client import ClaudeClient
from .config import CV_OUTPUT_DIR


def cv_to_structured_json(base_cv_text: str) -> dict[str, Any]:
    lines = [line.strip() for line in base_cv_text.splitlines() if line.strip()]
    return {
        "summary": lines[0] if lines else "",
        "skills": [l for l in lines if "," in l][:10],
        "experience": lines[1:40],
        "full_text": base_cv_text[:12000],
    }


async def tailor_cv(
    *,
    claude: ClaudeClient,
    base_cv_path: Path,
    base_cv_text: str,
    compressed_jd: dict[str, Any],
    company: str,
    job_id: int | None = None,
) -> tuple[str, float]:
    base_json = cv_to_structured_json(base_cv_text)
    diff = await claude.tailor_cv_gap(base_json, compressed_jd, job_id=job_id)

    output_docx = CV_OUTPUT_DIR / f"{_safe(company)}_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.docx"
    score = float(diff.get("match_score", 0))

    if Document and base_cv_path.exists() and base_cv_path.suffix.lower() == ".docx":
        doc = Document(str(base_cv_path))
        if doc.paragraphs:
            summary = diff.get("summary_rewrite")
            if summary:
                doc.paragraphs[0].text = str(summary)
        if diff.get("skills_add"):
            doc.add_paragraph("Additional Relevant Skills: " + ", ".join(diff["skills_add"][:8]))
        keywords = diff.get("keywords_to_inject", [])
        if keywords:
            doc.add_paragraph("Keywords: " + ", ".join(keywords[:12]))
        doc.save(str(output_docx))
    else:
        if base_cv_path.exists():
            shutil.copy(base_cv_path, output_docx)
        else:
            output_docx.write_text(json.dumps(diff, indent=2), encoding="utf-8")

    output_pdf = output_docx.with_suffix(".pdf")
    _try_convert_to_pdf(output_docx, output_pdf)

    final_score = _python_score_check(base_cv_text + " " + " ".join(diff.get("keywords_to_inject", [])), compressed_jd)
    return str(output_pdf if output_pdf.exists() else output_docx), max(score, final_score)


def _python_score_check(cv_text: str, compressed_jd: dict[str, Any]) -> float:
    required = compressed_jd.get("required_skills", [])
    if not required:
        return 80.0
    low_cv = cv_text.lower()
    hits = sum(1 for req in required if str(req).lower()[:40] in low_cv)
    return round(100 * hits / max(1, len(required)), 2)


def _try_convert_to_pdf(docx_path: Path, pdf_path: Path) -> None:
    try:
        subprocess.run(
            [
                "libreoffice",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(docx_path.parent),
                str(docx_path),
            ],
            check=False,
            capture_output=True,
            text=True,
        )
    except FileNotFoundError:
        return
    if pdf_path.exists():
        return


def _safe(name: str) -> str:
    return "".join(ch if ch.isalnum() or ch in "-_" else "_" for ch in name.lower()).strip("_") or "company"
